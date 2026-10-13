"""Fill the upGrad report starter (.docx) with results and charts.

Charts are the PNG outputs extracted from the executed notebook; every
number comes from that notebook's saved outputs.

Usage:
    python tools/build_report.py <starter.docx> <images_dir> <output.docx>
"""

import os
import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

TITLE_LINE = "EDA_Optimising_NYC_Taxis_Ishan_Ghosh"

INTRO = [
    "**Problem statement.** As an analyst for a new taxi operation in New York City, the goal is to use the 2023 yellow taxi trip records to find patterns that improve service efficiency, maximise revenue and improve the passenger experience.",
    "**Approach.** The twelve monthly files (38.3 million trips) were sampled by date and hour, cleaned (column fixes, missing values, outliers), and analysed in two stages: general EDA on time, money and location, then detailed EDA on operations, pricing and customer behaviour. All counts and revenue are scaled from the sample to full-year estimates using sample weights.",
]

ASSUMPTIONS = [
    "**Sampling:** 0.75% of the trips in every date-hour; each sampled trip represents about 133 real trips (its sample weight).",
    "**Reversal records:** rows with a negative fare or total are refunds, disputes or voids that cancel an earlier charge, so they were dropped instead of being made positive.",
    "**Missing values:** filled with the most likely value: passenger count and rate code with the mode (1), congestion surcharge with the most common value for the pickup zone, and the airport fee only for JFK/LaGuardia pickups.",
    "**Outliers:** removed only when physically impossible or clearly a recording error (over 6 passengers, over 250 miles, over 100 mph, zero/negative duration, zero distance and fare between different zones).",
    "**Night hours:** 11 PM to 5 AM.",
    "**Fare per mile:** uses trips of at least 0.5 miles, because the fixed starting fare makes very short trips look extremely expensive per mile.",
    "**Tips:** analysed on credit-card trips only, because cash tips are not recorded. Tip % = tip ÷ fare × 100.",
    "**Speed analysis:** leaves out trips that start and end in the same zone and trips longer than 180 minutes; route averages need at least 10 sampled trips.",
]

# anchor (start of the starter's task line) -> list of content items:
#   ("p", text)   paragraph, **bold** supported
#   ("img", file, width_in, caption)
CONTENT = {
    "Loading the dataset": [
        ("p", "Each monthly parquet file has about 3 million rows (January: 3,066,766; 19 columns), 38.3 million for the year. Loading and combining everything at once is impractical, so each file was read one at a time and sampled."),
    ],
    "Sample the data and combine the files": [
        ("p", "Within each file, trips were grouped by **pickup date and hour**, and 0.75% of every group was sampled with `groupby().sample()`. Sampling every date-hour keeps the daily, weekly and seasonal shape of demand exactly, and avoids one date or hour being over- or under-represented by chance."),
        ("p", "Each trip carries a **sample weight** (trips in its date-hour ÷ trips sampled from it), used to scale counts and revenue back up. Result: **287,296 sampled trips** (within the suggested 250k–300k), saved as a parquet file. The weights add up to 38,309,324 trips vs 38,310,226 in the files (99.998%), confirming the sample represents the full year."),
    ],
    "Fix the index": [
        ("p", "The index was reset after combining the months. `store_and_fwd_flag` was dropped: it only records whether a trip record was buffered in the vehicle before transmission and is not used in any analysis."),
    ],
    "Combine the two airport_fee columns": [
        ("p", "January names the column `airport_fee`, other months `Airport_fee`. No trip had a value in both, so they were combined with `combine_first` into one `airport_fee` column."),
        ("p", "**Negative monetary values:** 2,833 trips had a negative fare. They were not linked to a rate code (88% were standard rate, like the data overall) but to the payment type: Dispute 1,679, Cash 665, No charge 444. In 97.2% of them every charge was negative, mirroring a normal fare. These are reversal records, so the 2,854 rows with a negative fare or total (1.0% of the sample) were dropped. One trip with a negative `extra` remained and was kept and flagged."),
    ],
    "Find the proportion of missing values in each column": [
        ("p", "Four columns were each missing **3.45%** of values: `passenger_count`, `RatecodeID`, `congestion_surcharge` and `airport_fee`. All 9,804 missing values fall on **the same rows**, and every one of those rows has `payment_type` 0, a code not in the data dictionary. These records come from a source that does not fill the optional fields; their fares and locations are valid, so they were kept."),
    ],
    "Handling missing values in passenger_count": [
        ("p", "Missing counts were filled with the mode (1 passenger, 77% of trips). The 4,324 trips with 0 passengers were also set to 1: a completed, paid trip carried at least one passenger, so 0 is a missed driver entry."),
    ],
    "Handle missing values in RatecodeID": [
        ("p", "1,583 trips had `RatecodeID` 99, which is not a valid code, so they were treated as missing too. Missing codes were filled with the mode, 1 (Standard rate, about 95% of known codes)."),
    ],
    "Impute NaN in congestion_surcharge": [
        ("p", "The surcharge applies to trips in Manhattan south of 96th Street, so it depends on location. Missing values were filled with the most common surcharge for the same pickup zone (overall mode as a fallback). The missing `airport_fee` values were set to 1.75 USD for JFK/LaGuardia pickups and 0 elsewhere. No missing values remain."),
    ],
    "Check outliers in payment type, trip distance and tip amount columns": [
        ("img", "c064_01.png", 6.5, "Box plots before outlier handling: a few extreme distances, fares and tips stretch each axis."),
        ("p", "**Payment type:** 9,771 trips had `payment_type` 0 (not in the dictionary). They are real trips with fares, so the code was replaced with 5 (Unknown) instead of dropping them."),
        ("p", "**Trip distance:** 6 trips were over 250 miles (up to 44,409 miles, a meter or GPS error), 16 had zero distance and zero fare between different zones, and 1 had almost no distance but a fare over 300 USD. All were dropped."),
        ("p", "**Tip amount:** tips reach 200 USD, but no tip is negative after removing reversals and high tips are plausible, so tips were kept. 22.9% of tips are zero, mostly cash trips, where tips are not recorded."),
        ("p", "**Other:** 7 trips with over 6 passengers, 122 trips with zero or negative duration and 140 trips faster than 100 mph were dropped. In total only 274 rows (0.1%) were removed. The cleaned data has **284,168 trips**."),
    ],
    "Classify variables into categorical and numerical": [
        ("p", "**Categorical:** VendorID, RatecodeID, PULocationID, DOLocationID, payment_type (nominal codes, so the numbers are labels, not quantities) and pickup_hour (ordinal)."),
        ("p", "**Numerical:** passenger_count (discrete), trip_distance and trip_duration (continuous), and all monetary columns: fare_amount, extra, mta_tax, tip_amount, tolls_amount, improvement_surcharge, total_amount, congestion_surcharge, airport_fee."),
        ("p", "**Datetime:** tpep_pickup_datetime and tpep_dropoff_datetime, used to derive hour, day, month and quarter."),
    ],
    "Analyse the distribution of taxi pickups by hours, days of the week, and months": [
        ("img", "c079_02.png", 5.8, "Estimated pickups by hour of day."),
        ("p", "Demand is lowest at 4–5 AM, rises steeply from 6 AM and peaks at **6 PM (2.68 million trips a year)**. The evening (3–7 PM) is the busiest part of the day."),
        ("img", "c080_03.png", 5.8, "Average pickups per calendar day, by day of week."),
        ("p", "**Thursday** is the busiest day (114,208 trips per day) followed by Wednesday; Sunday (90,336) and Monday (90,787) are the quietest. Averages per calendar day are used because 2023 has 53 Sundays."),
        ("img", "c081_04.png", 5.8, "Estimated pickups by month."),
        ("p", "May and October are the busiest months (~3.48 million trips each). Demand dips in summer: July–September average about 90,000–94,000 trips a day, against 112,000 in May and October."),
    ],
    "Filter out the zero/negative values in fares, distance and tips": [
        ("p", "After cleaning there are no negative values. Zeros remain: fare 68 trips, total 36, distance 5,453 (1.9%, of which 2,209 start and end in the same zone) and tip 65,109 (22.9%)."),
        ("p", "A copy without zero fares, totals and distances (**278,684 trips**) is used for distance-based measures such as fare per mile and speed. Zero tips were kept: 73% of them are cash trips, where tips are not recorded, and removing real non-tippers would overstate tipping. Zero-distance trips stay in the main data for demand and revenue counts."),
    ],
    "Analyse the monthly revenue trends": [
        ("img", "c091_05.png", 6.0, "Estimated monthly revenue (bars) and average revenue per trip (line)."),
        ("p", "Estimated 2023 revenue is about **1.10 billion USD**. Monthly revenue follows demand: October is highest (103.9M USD) and February lowest (79.1M USD). Revenue per trip rose steadily from 27.56 USD in January to 30.25 USD in September, so revenue grew faster than trips."),
    ],
    "Find the proportion of each quarter": [
        ("img", "c093_06.png", 4.8, "Share of yearly revenue by quarter."),
        ("p", "Q4 (27.0%) and Q2 (26.7%) bring in the most revenue; Q3 (22.9%) the least, reflecting the summer dip. Q1 is 23.5%."),
    ],
    "Analyse and visualise the relationship between distance and fare amount": [
        ("img", "c095_07.png", 5.5, "Trip distance vs fare (random sample of 30,000 trips)."),
        ("p", "Fare rises almost linearly with distance (correlation **0.945**). The orange horizontal line at 70 USD is the JFK flat fare, which does not depend on distance."),
    ],
    "Analyse the relationship between fare/tips and trips/passengers": [
        ("img", "c097_08.png", 5.2, "Trip duration vs fare."),
        ("p", "**Fare vs duration:** weaker correlation (0.275). Time spent in slow traffic adds little to the fare compared with distance, and flat-fare airport trips vary widely in duration."),
        ("img", "c098_09.png", 5.2, "Fare by passenger count (outliers hidden)."),
        ("p", "**Fare vs passengers:** almost no relationship (0.042). Median fares are 13.5–14.9 USD for every group size; groups of 2–4 take slightly longer trips, so their averages are a little higher."),
        ("img", "c099_10.png", 5.2, "Trip distance vs tip, credit-card trips."),
        ("p", "**Tip vs distance:** tips rise with distance: correlation 0.592 on all trips and **0.734 on card trips** (cash tips are recorded as 0)."),
    ],
    "Analyse the distribution of different payment types": [
        ("img", "c101_11.png", 5.5, "Share of trips by payment type."),
        ("p", "Credit card dominates (**78.7%**), followed by cash (16.7%). Unknown (the recoded type 0) is 3.4%, disputes 0.7% and no-charge trips 0.5%. Card trips average a 4.40 USD tip, cash trips 0, because cash tips are not recorded."),
    ],
    "Load the taxi zones shapefile and display it": [
        ("img", "c110_12.png", 4.3, "NYC taxi zones coloured by borough."),
        ("p", "The shapefile has 263 zones, each with a unique `LocationID`, a zone name, a borough and its shape."),
    ],
    "Merge the zone data with trips data": [
        ("p", "Zone and borough names were merged onto every trip twice, once using the pickup ID and once using the dropoff ID. The row count stayed the same (284,168). 2,764 pickups use IDs 264/265 (Unknown / Outside NYC), which have no shape; they are labelled as such."),
    ],
    "Find the number of trips for each zone/location ID": [
        ("p", "Grouping by pickup zone gives the trips per zone. The busiest pickup zones are JFK Airport (1.97 million trips a year), Upper East Side South (1.77M), Midtown Center (1.76M), Upper East Side North (1.60M), Midtown East (1.34M), Penn Station/Madison Sq West (1.30M) and LaGuardia Airport (1.26M)."),
    ],
    "Add the number of trips for each zone to the zones dataframe": [
        ("p", "The trip counts were merged into the zones GeoDataFrame on `LocationID`, with zones that had no pickups set to 0."),
    ],
    "Plot a map of the zones showing number of trips": [
        ("img", "c121_13.png", 4.8, "Estimated pickups by zone, 2023."),
        ("p", "Pickups are concentrated in Manhattan (**89.3%** of all pickups) and at the two airports in Queens (9.8% for the borough). Brooklyn (0.7%), the Bronx and Staten Island have very little yellow-taxi demand."),
    ],
    "Conclude with results": [
        ("p", "**When:** weekday evenings (peak 6 PM), Thursdays, and May and October are busiest; early mornings, Sundays/Mondays and summer are quietest."),
        ("p", "**Revenue:** about 1.10 billion USD, strongest in Q2 and Q4. Revenue per trip increased during the year."),
        ("p", "**Fares:** driven by distance (correlation 0.945), much less by duration (0.275) and not by passengers (0.042). Tips grow with distance (0.734 on card trips)."),
        ("p", "**Where:** Manhattan (Upper East Side, Midtown) and the airports (JFK, LaGuardia) generate almost all trips."),
    ],
    "Identify slow routes by comparing average speeds on different routes": [
        ("img", "c131_14.png", 6.3, "Median trip speed by day and hour (red = slow)."),
        ("p", "Speed on a route in an hour = average distance ÷ average duration, for routes with at least 10 sampled trips. The slowest routes are short **Midtown** hops: Midtown East → Times Square at 1 PM (2.95 mph, 0.9 miles in 19 minutes), and Penn Station → Garment District at 9 AM, 10 AM, 3 PM and 6 PM (about 3.2–3.5 mph). From 8 AM to 11 PM the slowest route in every hour is inside Midtown, at 3–4.3 mph."),
        ("p", "The heatmap shows weekday daytime (8 AM–7 PM) is the slowest period across the city, and early mornings the fastest."),
    ],
    "Calculate the hourly number of trips and identify the busy hours": [
        ("img", "c135_15.png", 5.8, "Sampled trips by pickup hour."),
        ("p", "The busiest hour is **6 PM** with 20,087 sampled trips, followed by 5 PM, 7 PM, 4 PM and 3 PM."),
    ],
    "Scale up the number of trips from above to find the actual number of trips": [
        ("p", "Dividing by the sampling fraction (0.0075) gives the actual trips in the five busiest hours of 2023: **6 PM 2.68 million**, 5 PM 2.55M, 7 PM 2.39M, 4 PM 2.35M and 3 PM 2.34M. The scaled totals agree with the sample-weight estimates to within 0.1%. At 6 PM that is about 7,340 trips every day."),
    ],
    "Compare hourly traffic on weekdays and weekends": [
        ("img", "c140_16.png", 5.8, "Average hourly pickups on a weekday vs a weekend day."),
        ("img", "c140_17.png", 6.3, "Average pickups by day of week and hour."),
        ("p", "**Weekdays** follow a commuter pattern: a morning rise from 6 AM (3,477 trips a day at 7 AM) and a peak at **6 PM with 7,833 trips a day**. **Weekends** start slowly (1,244 at 7 AM) but are far busier late at night: between midnight and 4 AM there are 2.5–5.8 times more trips than on a weekday night (3,122 vs 565 at 2 AM). Tuesday–Thursday at 5–6 PM are the busiest slots of the week, and Friday and Saturday stay busy past midnight."),
        ("p", "Knowing these patterns lets the operator schedule shifts to demand, plan breaks and maintenance in the quiet hours, and move cabs to nightlife areas before the weekend late-night peak."),
    ],
    "Identify the top 10 zones with high hourly pickups and drops": [
        ("img", "c143_18.png", 6.3, "Top 10 pickup and dropoff zones (average trips per hour)."),
        ("p", "**Pickups:** JFK Airport (224.6 an hour), Upper East Side South (201.8), Midtown Center (201.0), Upper East Side North (182.4), Midtown East (153.5), Penn Station (148.3), LaGuardia (144.2), Lincoln Square East, Times Square and Murray Hill. **Dropoffs:** Upper East Side North (189.5), Upper East Side South (181.8), Midtown Center (169.1), Times Square, Murray Hill, Midtown East, Upper West Side South, Lincoln Square East, Lenox Hill West and East Chelsea."),
        ("img", "c143_19.png", 6.3, "Hourly pickups in the top 10 pickup zones."),
        ("img", "c143_20.png", 6.3, "Hourly dropoffs in the top 10 dropoff zones (by dropoff hour)."),
        ("p", "Midtown Center shows the commuter cycle: its dropoffs peak at 8–9 AM (people arriving for work) and its pickups at 5–6 PM (people leaving). The Upper East Side zones are busy from late morning to early evening for both pickups and dropoffs, and JFK pickups stay high from 2 PM until 11 PM, peaking around 10 PM."),
    ],
    "Find the ratio of pickups and dropoffs in each zone": [
        ("p", "**Highest pickup/dropoff ratios** (cabs leave full more often than they arrive): East Elmhurst 8.25 (next to LaGuardia), JFK Airport 4.29, LaGuardia Airport 2.51, Penn Station 1.57, Greenwich Village South 1.38, West Village 1.38, Central Park 1.32, Midtown East 1.25."),
        ("p", "**Lowest ratios** (cabs arrive and leave empty): Columbia Street 0.01, Murray Hill-Queens 0.04, Roosevelt Island 0.04, Riverdale 0.05, **Newark Airport 0.06**, Middle Village 0.06, Forest Hills 0.08, Inwood 0.08: mostly residential outer-borough zones."),
    ],
    "Identify the top zones with high traffic during night hours": [
        ("img", "c147_21.png", 6.3, "Top 10 night pickup (by pickup time) and dropoff (by dropoff time) zones, 11 PM–5 AM."),
        ("p", "Night **pickups** are led by nightlife areas: East Village (323,000 trips a year), JFK Airport (265,000), West Village (261,000), Clinton East, Lower East Side, Greenwich Village South and Times Square. Night **dropoffs** are mainly residential Manhattan: East Village, Clinton East, East Chelsea, Murray Hill, Gramercy, Lenox Hill West and Yorkville West."),
    ],
    "Find the revenue share for nighttime and daytime hours": [
        ("img", "c149_22.png", 4.8, "Share of trips and revenue, day vs night."),
        ("p", "Night (11 PM–5 AM) brings in **11.5% of revenue** from 11.5% of trips; day brings in 88.5%. Night trips are longer on average (3.81 vs 3.41 miles) but earn the same per trip (28.99 vs 28.98 USD), so revenue per mile is lower at night."),
    ],
    "For the different passenger counts, find the average fare per mile per passenger": [
        ("img", "c152_23.png", 5.5, "Average fare per mile per passenger."),
        ("p", "Average fare per mile is almost identical for every group size (7.36–7.57 USD). Divided by passengers, it falls from **7.52 USD per mile for one passenger to 1.24 USD for six**: group trips are far cheaper per person."),
    ],
    "Find the average fare per mile by hours of the day and by days of the week": [
        ("img", "c154_24.png", 6.3, "Average fare per mile by hour and by day of week."),
        ("p", "Fare per mile is highest in the middle of the day (**8.20 USD at noon**, above 8 USD from 11 AM to 5 PM), when slow traffic adds time charges, and lowest at 4–5 AM (5.7 USD), when roads are clear and trips are longer. By day, Wednesday (7.88) and Thursday (7.80) are highest and Sunday (6.81) lowest."),
    ],
    "Analyse the average fare per mile for the different vendors": [
        ("img", "c156_25.png", 5.8, "Average fare per mile by vendor and hour."),
        ("p", "Creative Mobile Technologies averages 7.56 USD per mile (median 7.12) and VeriFone 7.47 (median 7.02), with the same hourly pattern. Myle Technologies has only 50 sampled trips, too few for reliable conclusions."),
    ],
    "Compare the fare rates of different vendors in a distance-tiered fashion": [
        ("img", "c158_26.png", 5.5, "Average fare per mile by distance tier and vendor."),
        ("p", "Per-mile fares fall with distance because the fixed starting charge is spread over more miles: **up to 2 miles** CMT 9.03 vs VeriFone 9.08 USD; **2–5 miles** 6.40 vs 6.54; **over 5 miles** 4.43 vs 4.49. The two main vendors are within 0.15 USD per mile in every tier, so pricing is effectively matched (Myle's 28.47 in the shortest tier comes from very few trips)."),
    ],
    "Analyse the tip percentages": [
        ("img", "c161_27.png", 6.5, "Average tip % by distance tier, passenger count and pickup hour (card trips)."),
        ("p", "Card passengers tip **24.9%** of the fare on average (median 25.9%). Tip % falls with distance: 27.7% up to 2 miles, 22.2% for 2–5 miles and 20.4% above 5 miles, while the share of low tips (<10%) rises from 7.1% to 15.2%. Passenger count makes little difference (24.3–25.4%). Tips are highest in the evening (about 26.4% at 6–7 PM) and lowest early in the morning, when about 17% of trips tip under 10% (3–6 AM)."),
        ("p", "Compared with high tippers (>25%), low-tip trips are twice as long (4.84 vs 2.31 miles; 25 vs 13 minutes) and more expensive (25.65 vs 14.43 USD), more often at night (15% vs 11%) and more often airport-rate trips (6% vs 1%). Passengers seem to tip a smaller percentage on large fares."),
    ],
    "Analyse the trends in passenger count": [
        ("img", "c165_28.png", 6.3, "Average passengers per trip by hour and by day."),
        ("img", "c165_29.png", 6.3, "Average passengers per trip by day and hour."),
        ("p", "77.4% of trips carry one passenger and 14.8% two; only 7.8% carry three or more. Group sizes are largest on weekends and late at night (social trips) and smallest on weekday mornings (commuters travelling alone)."),
    ],
    "Analyse the variation of passenger counts across zones": [
        ("img", "c168_30.png", 4.8, "Average passengers per pickup by zone (zones with at least 50 sampled trips)."),
        ("p", "The highest average group sizes are in tourist and airport zones: Battery Park (1.66), South Ozone Park (1.57), Seaport (1.54), Financial District South (1.52), JFK Airport (1.51) and Times Square (1.48), where around 30% of trips carry two or more people. The lowest are residential outer-borough zones such as Crown Heights North (1.10), Woodside and Long Island City (1.16–1.17)."),
    ],
    "Analyse the pickup/dropoff zones or times when extra charges are applied more frequently": [
        ("img", "c171_31.png", 5.5, "How often each charge is applied."),
        ("img", "c171_32.png", 6.0, "Share of trips charged, by pickup hour."),
        ("p", "The improvement surcharge (99.96%) and MTA tax (99.1%) apply to almost every trip and the congestion surcharge to 92.2%. The `extra` charge applies to 59.9% of trips, the airport fee to 8.5% and tolls to 8.1%."),
        ("p", "**By time:** `extra` is charged on most trips overnight and in the 4–8 PM rush hour but only about a third of trips between 6 AM and 3 PM, matching the overnight and rush-hour surcharges. It applies to 60–66% of weekday trips but only about 50% at weekends (no weekday rush-hour charge)."),
        ("p", "**By zone:** `extra` is highest for LaGuardia pickups (98.7%) and nightlife zones (Lower East Side 78%, Greenwich Village South 77%, East Village 73%). The congestion surcharge applies to almost all trips in Manhattan zones (about 99.7–99.9%). The airport fee is charged on 98.3% of LaGuardia and 93.5% of JFK pickups, and tolls on 93.3% of trips to Newark Airport."),
    ],
    "Recommendations to optimize routing and dispatching": [
        ("p", "**Shape the fleet to the daily demand curve.** On weekdays, bring cabs online from 6 AM and have the most cabs on the road from 3 PM to 7 PM (6 PM averages about 7,800 trips per weekday). Keep the smallest fleet at 4–5 AM, when demand is under a tenth of the peak."),
        ("p", "**Run a separate weekend schedule.** Move driver hours from weekend mornings (7–8 AM is well under half the weekday level) to 11 PM–4 AM, when weekend demand is 2.5–5.8 times that of a weekday night."),
        ("p", "**Avoid Midtown cross-traffic in the daytime.** The slowest routes are short Midtown hops at 3–4 mph from 8 AM to 7 PM. Dispatch should assign the nearest cab on the same side of Midtown, avoid sending empty cabs through Midtown to reach a pickup, and show realistic ETAs on these routes."),
        ("p", "**Run the airports as dedicated queues.** JFK is the busiest pickup zone (about 225 pickups an hour) and both airports have far more pickups than dropoffs (JFK 4.3×, LaGuardia 2.5×), so they need a steady supply of cabs fed from nearby Queens dropoffs."),
        ("p", "**Plan for the seasons.** Daily demand is lowest in July–September and highest in May and October: schedule leave and maintenance in summer and run a full fleet in spring and Q4."),
    ],
    "Suggestions on strategically positioning cabs": [
        ("p", "**Daytime core (weekdays 7 AM–8 PM):** Upper East Side South and North, Midtown Center and East, Penn Station, Lincoln Square East, Times Square and Murray Hill, each with 125–200 pickups an hour on average."),
        ("p", "**Airports all day:** JFK and LaGuardia together take about 8.5% of all pickups; the neighbouring East Elmhurst zone has 8× more pickups than dropoffs."),
        ("p", "**Nightlife at night, especially Friday and Saturday:** East Village, West Village, Clinton East, Lower East Side, Greenwich Village South and Times Square for pickups. After night dropoffs in residential Manhattan (Murray Hill, Gramercy, Lenox Hill, Yorkville), cabs should head back to these zones."),
        ("p", "**Rebalance from zones where cabs end up empty:** outer-borough residential zones and Newark Airport have pickup/dropoff ratios below 0.1. After a drop there, guide drivers back to Manhattan, or to JFK/LaGuardia from Queens."),
        ("p", "**Larger vehicles where groups travel:** Battery Park, the Seaport, the Financial District, JFK and Times Square (about 30% of trips with 2+ passengers), and generally at night and at weekends."),
    ],
    "Propose data-driven adjustments to the pricing strategy": [
        ("p", "**Charge for extra passengers.** Fare per mile is the same (about 7.5 USD) whatever the group size, so a six-person trip costs 1.24 USD per mile per person. A small fee per extra passenger, or a group fare for larger vehicles, adds revenue without affecting the 77% of single-passenger trips."),
        ("p", "**Stay within the competitors' band.** The two main vendors are within 0.15 USD per mile in every distance tier (about 9.0, 6.5 and 4.5 USD per mile). A new operator should price within this band and compete on availability and service."),
        ("p", "**Lift the long-trip rate slightly.** Trips over 5 miles earn about 4.5 USD per mile, half the short-trip rate, and also get the lowest tips (20.4%). A modestly higher per-mile rate beyond 5 miles (outside the fixed airport fares) would improve earnings on the trips that tie up a cab longest."),
        ("p", "**Use off-peak incentives instead of higher peak fares.** Fare per mile peaks at midday (8.2 USD) and is lowest at 4 AM (5.7 USD), and Q3 has the lowest revenue share (22.9%). Discounts in summer and on weekday early mornings can fill idle capacity while peak prices stay competitive."),
        ("p", "**Make card tipping easy.** Card trips (78.7%) tip about 25% of the fare. Low tips are most common on long, airport-rate and 3–6 AM trips; suggested tip percentages on the payment screen help keep tips up."),
    ],
}


def add_text(paragraph, text):
    """Add text to a paragraph, rendering **bold** and `code` segments."""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def insert_after(paragraph, text=None, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.style = style or paragraph.style
    if text:
        add_text(new_paragraph, text)
    return new_paragraph


def set_task_line(paragraph):
    """Keep the starter's task text as a bold line, without its empty line breaks."""
    text = paragraph.text.strip()
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    run = paragraph.add_run(text)
    run.bold = True
    paragraph.paragraph_format.keep_with_next = True


def add_items(anchor, items, images_dir, body_style):
    last = anchor
    for item in items:
        if item[0] == "p":
            last = insert_after(last, item[1], body_style)
            last.paragraph_format.space_after = Pt(6)
        else:
            _, name, width, caption = item
            picture = insert_after(last, style=body_style)
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.keep_with_next = True  # keep the caption with the chart
            picture.add_run().add_picture(os.path.join(images_dir, name), width=Inches(width))
            last = insert_after(picture, style=body_style)
            last.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = last.add_run(caption)
            run.italic = True
            run.font.size = Pt(9)
            last.paragraph_format.space_after = Pt(10)
    return last


def build(starter_path, images_dir, output_path):
    document = docx.Document(starter_path)
    body_style = document.styles["normal"]

    # Remove the starter's empty placeholder paragraphs
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip() and not paragraph._p.xpath(".//w:drawing"):
            paragraph._p.getparent().remove(paragraph._p)

    paragraphs = document.paragraphs
    title, instructions = paragraphs[0], paragraphs[1]

    # Title line with the required name format, then intro and assumptions
    for run in list(instructions.runs):
        run._r.getparent().remove(run._r)
    name_run = instructions.add_run(TITLE_LINE)
    name_run.bold = True
    last = instructions
    for text in INTRO:
        last = insert_after(last, text, body_style)
    heading = insert_after(last, "Assumptions", document.styles["Heading 2"])
    last = heading
    for text in ASSUMPTIONS:
        last = insert_after(last, text, body_style)

    used = set()
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if paragraph.style.name != "normal":
            continue
        for anchor, items in CONTENT.items():
            if text.startswith(anchor) and anchor not in used:
                used.add(anchor)
                set_task_line(paragraph)
                add_items(paragraph, items, images_dir, body_style)
                break

    missing = set(CONTENT) - used
    assert not missing, f"Anchors not found in the starter: {missing}"
    document.save(output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3])
